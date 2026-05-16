from docx import Document

class ParsePRD:
    def parse_doc_file(file_path):
        try:
            if not file_path.lower().endswith('.docx'):
                print(f"不是docx文件：{file_path}")
                return None
            print(f"正在解析docx文件：{file_path}")
            # 读取文档
            doc = Document(file_path)

            # 提取所有的段落文本
            paragraphs = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)

            # print(paragraphs)

            # print("-=========================")
            tables_content = []
            for table in doc.tables:
                table_text = []
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        table_text.append("|".join(row_text))
                if table_text:
                    tables_content.append("表格:\n" + "\n".join(table_text))

            # print(tables_content)
            # print("-=========================")
            # 组合所有内容
            all_content = []

            if paragraphs:
                # 将段落列表转换为字符串
                paragraphs_str = "\n".join(paragraphs)
                all_content.append("文档内容：")
                all_content.append(paragraphs_str)
            if tables_content:
                # 将表格内容列表转换为字符串
                tables_str = "\n".join(tables_content)
                all_content.append("表格内容")
                all_content.append(tables_str)

            if not all_content:
                print("文档内容为空")
                return None
            result = "\n".join(all_content)
            # print(f"成功解析文档，提取{len(paragraphs)}个段落，{len(tables_content)}个表格")
            return result
        except Exception as e:
            print(f"解析文档失败：{e}")
            return None

# path = f"/Users/dongyueqian/PycharmProjects/case_agent/产品需求文档.docx"
# prd = ParsePRD.parse_doc_file(path)
# print(len(prd))